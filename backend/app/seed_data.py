import sqlite3
import os
import re
import hashlib
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from database import get_db, init_db
import sys

sys.stdout.reconfigure(encoding='utf-8')

EXCEL_SOURCE_PATH = r'C:\Users\Mada\Downloads\Mounir System\application.xlsx'
OUTPUT_DIR = r'C:\Users\Mada\Downloads\Mounir System'

def clean_phone(val):
    if val is None:
        return '01000000000'
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    s = re.sub(r'[^\d+]', '', s)
    if not s:
        return '01000000000'
    if len(s) == 10 and s.startswith('1'):
        s = '0' + s
    return s

def parse_age(val):
    if val is None:
        return 10
    s = str(val).replace('٫', '.').strip()
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except Exception:
        return 10

def parse_credits(val):
    if val is None:
        return 4
    try:
        return int(float(val))
    except Exception:
        return 4

def generate_student_password(st_code):
    # Generates a deterministic 6-char password e.g. Mn4829
    hash_val = int(hashlib.sha256(f"Mounir_{st_code}".encode()).hexdigest(), 16)
    num = (hash_val % 9000) + 1000
    return f"Mn{num}"

def generate_teacher_password(t_id):
    hash_val = int(hashlib.sha256(f"Teacher_{t_id}".encode()).hexdigest(), 16)
    num = (hash_val % 9000) + 1000
    return f"Tch@{num}"

def seed_all():
    print("Starting system re-seeding from application.xlsx...")
    init_db()
    conn = get_db()
    cursor = conn.cursor()

    # 1. Clear all previous data
    cursor.execute("DELETE FROM teachers")
    cursor.execute("DELETE FROM students")
    cursor.execute("DELETE FROM courses")
    cursor.execute("DELETE FROM enrollments")
    cursor.execute("DELETE FROM lectures")
    cursor.execute("DELETE FROM attendance")
    cursor.execute("DELETE FROM quizzes")
    cursor.execute("DELETE FROM quiz_questions")
    cursor.execute("DELETE FROM quiz_submissions")
    cursor.execute("DELETE FROM support_tickets")
    cursor.execute("DELETE FROM notifications")
    cursor.execute("DELETE FROM payments")
    cursor.execute("DELETE FROM teacher_payouts")
    cursor.execute("DELETE FROM users")

    # 2. Parse application.xlsx
    wb = openpyxl.load_workbook(EXCEL_SOURCE_PATH)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))[1:]

    students_dict = {}  # code -> dict
    teachers_set = set()
    courses_dict = {}   # course_name -> dict
    enrollments_list = []

    for idx, r in enumerate(rows, start=2):
        if not any(r):
            continue
        g_id, st_code, st_name, age, phone, status, remaining, t_name, g_name, time_val, days = r[:11]

        st_code = str(st_code or '').strip()
        st_name = str(st_name or '').strip()
        t_name = str(t_name or 'معلم عام').strip()
        g_name = str(g_name or 'مجموعة عامة').strip()

        if not st_code or not st_name:
            continue

        teachers_set.add(t_name)

        if g_name not in courses_dict:
            courses_dict[g_name] = {
                'group_id': str(g_id or '').strip(),
                'name': g_name,
                'teacher': t_name,
                'days': str(days or '').strip(),
                'time': str(time_val or '').strip()
            }

        if st_code not in students_dict:
            students_dict[st_code] = {
                'code': st_code,
                'name': st_name,
                'age': parse_age(age),
                'phone': clean_phone(phone),
                'status': 'active' if str(status or '').strip() == 'نشط' else 'frozen',
                'raw_status': str(status or '').strip()
            }

        enrollments_list.append({
            'student_code': st_code,
            'course_name': g_name,
            'teacher_name': t_name,
            'remaining_credits': parse_credits(remaining),
            'status': 'active' if str(status or '').strip() == 'نشط' else 'frozen',
            'raw_status': str(status or '').strip()
        })

    # 3. Insert Teachers
    sorted_teachers = sorted(list(teachers_set))
    teacher_name_to_id = {}
    teacher_credentials = []

    for t_id, t_name in enumerate(sorted_teachers, start=1):
        teacher_name_to_id[t_name] = t_id
        email = f"teacher_{t_id}@monir-academy.edu.eg"
        t_username = f"T{t_id:03d}"
        t_password = generate_teacher_password(t_id)

        cursor.execute('''
        INSERT INTO teachers (id, name, specialty, email, phone, bio, rate_per_session, rate_per_private_session, late_count, cancellation_count, active)
        VALUES (?, ?, 'مسار تحفيظ والقرآن والتجويد', ?, '01000000000', 'معلم معتمد في الأكاديمية', 150.0, 80.0, 0, 0, 1)
        ''', (t_id, t_name, email))

        # Teacher user login
        cursor.execute('''
        INSERT INTO users (username, password_hash, role, related_id, full_name, email, phone, status)
        VALUES (?, ?, 'teacher', ?, ?, ?, '01000000000', 'active')
        ''', (t_username, t_password, t_id, t_name, email))

        teacher_credentials.append({
            'id': t_id,
            'name': t_name,
            'username': t_username,
            'password': t_password,
            'email': email
        })

    # 4. Insert Admin Account
    cursor.execute('''
    INSERT INTO users (username, password_hash, role, related_id, full_name, email, phone, status)
    VALUES ('admin', 'admin2026', 'admin', NULL, 'إدارة أكاديمية منير', 'admin@monir-academy.edu.eg', '01000000000', 'active')
    ''')

    # 5. Insert Courses
    course_name_to_id = {}
    for c_id, (c_name, c_info) in enumerate(courses_dict.items(), start=1):
        course_name_to_id[c_name] = c_id
        cursor.execute('''
        INSERT INTO courses (id, name, track_name, target_age, min_age, max_age, max_capacity, default_duration_minutes, description, price_per_block, total_lectures, default_day, default_time)
        VALUES (?, ?, 'مسار القرآن والتدبر', '6 - 16 سنة', 6, 16, 10, 60, ?, 450.0, 8, ?, ?)
        ''', (c_id, c_name, f"مجموعة {c_name}", c_info['days'] or 'السبت', c_info['time'] or '18:00'))

    # 6. Insert Students & User Credentials
    student_code_to_id = {}
    student_credentials = []

    for s_id, (st_code, s_info) in enumerate(students_dict.items(), start=1):
        student_code_to_id[st_code] = s_id
        qr_url = f"https://api.qrserver.com/v1/create-qr-code/?size=150x150&data={st_code}"
        parent_name = f"ولي أمر {s_info['name']}"

        cursor.execute('''
        INSERT INTO students (id, name, student_code, age, phone, parent_name, parent_phone, qr_code)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ''', (s_id, s_info['name'], st_code, s_info['age'], s_info['phone'], parent_name, s_info['phone'], qr_url))

        # Student User Login
        username = st_code  # e.g. ST0001
        password = generate_student_password(st_code)

        cursor.execute('''
        INSERT INTO users (username, password_hash, role, related_id, full_name, phone, status)
        VALUES (?, ?, 'student', ?, ?, ?, ?)
        ''', (username, password, s_id, s_info['name'], s_info['phone'], s_info['status']))

        # Find first group name & teacher name for excel report
        st_enrolls = [e for e in enrollments_list if e['student_code'] == st_code]
        first_group = st_enrolls[0]['course_name'] if st_enrolls else ''
        first_teacher = st_enrolls[0]['teacher_name'] if st_enrolls else ''
        first_rem = st_enrolls[0]['remaining_credits'] if st_enrolls else 4

        student_credentials.append({
            'code': st_code,
            'name': s_info['name'],
            'username': username,
            'password': password,
            'age': s_info['age'],
            'phone': s_info['phone'],
            'group_name': first_group,
            'teacher_name': first_teacher,
            'remaining_credits': first_rem,
            'status': s_info['raw_status']
        })

    # 7. Insert Enrollments
    for e in enrollments_list:
        st_id = student_code_to_id.get(e['student_code'])
        t_id = teacher_name_to_id.get(e['teacher_name'], 1)
        if not st_id:
            continue

        cursor.execute('''
        INSERT OR IGNORE INTO enrollments (
            student_id, course_name, teacher_id, unlocked_blocks, total_lectures_unlocked,
            renewal_count, remaining_credits, excuse_count, max_allowed_excuses, current_surah, status
        ) VALUES (?, ?, ?, ?, 4, 0, ?, 0, 1, 'سورة الرحمن - المرحلة الأولى', ?)
        ''', (st_id, e['course_name'], t_id, 1 if e['remaining_credits'] > 0 else 0, e['remaining_credits'], e['status']))

    # 8. Seed Lectures per course (8 per course)
    lec_id = 1
    for c_name in courses_dict.keys():
        for lec_num in range(1, 9):
            block_num = 1 if lec_num <= 4 else 2
            status_str = 'completed' if lec_num <= 2 else ('live' if lec_num == 3 else 'scheduled')
            cursor.execute('''
            INSERT INTO lectures (
                id, course_name, lecture_number, title, description, block_number,
                scheduled_time, google_meet_url, status
            ) VALUES (?, ?, ?, ?, ?, ?, '2026-09-15 18:00', 'https://meet.google.com/mounir-class', ?)
            ''', (lec_id, c_name, lec_num, f"المحاضرة {lec_num}: متابعة الحفظ والتجويد", f"شرح وتسميع الحصة رقم {lec_num}", block_num, status_str))
            lec_id += 1

    conn.commit()
    conn.close()

    print(f"Database seeded successfully with {len(students_dict)} students, {len(sorted_teachers)} teachers, and {len(courses_dict)} courses!")

    # 9. Create Excel credentials file in C:\Users\Mada\Downloads\Mounir System
    export_students_excel(student_credentials)
    export_teachers_excel(teacher_credentials)

def export_students_excel(credentials):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "بيانات دخول الطلاب"
    ws.sheet_view.rightToLeft = True

    # Styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    headers = [
        "كود الطالب", "اسم الطالب", "اسم المستخدم", "كلمة المرور",
        "السن", "رقم التواصل", "اسم الجروب / الكورس", "اسم المعلم",
        "رصيد الحصص المتبقي", "حالة الحساب"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for r_idx, c in enumerate(credentials, start=2):
        row_data = [
            c['code'], c['name'], c['username'], c['password'],
            c['age'], c['phone'], c['group_name'], c['teacher_name'],
            c['remaining_credits'], c['status']
        ]
        ws.append(row_data)

        fill_to_use = alt_fill if r_idx % 2 == 1 else None
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border
            if fill_to_use:
                cell.fill = fill_to_use

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    file_path_1 = os.path.join(OUTPUT_DIR, 'students_credentials.xlsx')
    file_path_2 = os.path.join(OUTPUT_DIR, 'بيانات_دخول_الطلاب.xlsx')

    wb.save(file_path_1)
    wb.save(file_path_2)
    print(f"Saved Student Credentials Excel files:\n  - {file_path_1}\n  - {file_path_2}")

def export_teachers_excel(credentials):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "بيانات دخول المعلمين"
    ws.sheet_view.rightToLeft = True

    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid") # Teal Dark
    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = ["رقم المعلم", "اسم المعلم", "اسم المستخدم", "كلمة المرور", "البريد الإلكتروني"]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for r_idx, c in enumerate(credentials, start=2):
        row_data = [c['id'], c['name'], c['username'], c['password'], c['email']]
        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    file_path = os.path.join(OUTPUT_DIR, 'teachers_credentials.xlsx')
    wb.save(file_path)
    print(f"Saved Teacher Credentials Excel file:\n  - {file_path}")

if __name__ == '__main__':
    seed_all()
